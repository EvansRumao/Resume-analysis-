from pathlib import Path
from io import BytesIO
import html
import re
import pickle

import pymupdf
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

BASE_DIR = Path(__file__).resolve().parents[1]
MODEL_PATH = BASE_DIR / "model" / "resume_screening_model.pkl"

app = FastAPI(
    title="Resume Analyzer API",
    description="Resume category prediction and job-description matching API",
    version="1.0.0",
)

# Development CORS. Restrict allow_origins before deploying publicly.
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "https://resume-analysis-fv7y-quweevwdj-si-o3.vercel.app/"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

bundle = None
model = None
model_error = None

if MODEL_PATH.exists():
    try:
        # Security: only load a pickle file you created and trust.
        with MODEL_PATH.open("rb") as f:
            bundle = pickle.load(f)
        model = bundle.get("model") if isinstance(bundle, dict) else bundle
    except Exception as exc:
        model_error = f"Could not load model: {exc}"
else:
    model_error = (
        f"Model file not found at {MODEL_PATH}. "
        "Copy resume_screening_model.pkl into backend/model/."
    )

# Fallback list used by the matching endpoint; bundle skills take precedence.
DEFAULT_SKILLS = [
    "python", "java", "javascript", "typescript", "c++", "sql", "mysql",
    "postgresql", "mongodb", "firebase", "html", "css", "react", "node.js",
    "express", "spring boot", "fastapi", "django", "flask", "machine learning",
    "deep learning", "natural language processing", "nlp", "artificial intelligence",
    "generative ai", "data science", "pandas", "numpy", "scikit-learn",
    "tensorflow", "pytorch", "power bi", "tableau", "excel", "aws", "azure",
    "gcp", "docker", "kubernetes", "git", "github", "linux", "rest api",
    "microservices", "communication", "problem solving", "data analysis",
    "statistics", "cybersecurity", "network security", "ethical hacking",
    "penetration testing", "siem", "splunk", "risk management"
]
SKILLS = (bundle.get("skills", DEFAULT_SKILLS)
          if isinstance(bundle, dict) else DEFAULT_SKILLS)


class TextPredictionRequest(BaseModel):
    text: str = Field(min_length=1, max_length=200_000)
    top_k: int = Field(default=5, ge=1, le=20)


class MatchRequest(BaseModel):
    resume_text: str = Field(min_length=1, max_length=200_000)
    job_description: str = Field(min_length=1, max_length=50_000)


def clean_text(text: str) -> str:
    text = html.unescape(str(text)).lower()
    text = re.sub(r"https?://\S+|www\.\S+", " ", text)
    text = re.sub(r"\S+@\S+", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def extract_skills(text: str) -> list[str]:
    cleaned = clean_text(text)
    found = []
    for skill in SKILLS:
        pattern = r"(?<![a-z0-9])" + re.escape(skill.lower()) + r"(?![a-z0-9])"
        if re.search(pattern, cleaned):
            found.append(skill)
    return sorted(set(found))


def extract_resume_text(filename: str, content: bytes) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
           with pymupdf.open(stream=content, filetype="pdf") as pdf:
            text = "\n".join(page.get_text() for page in pdf)
        elif name.endswith(".docx"):
            doc = Document(BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        elif name.endswith(".txt"):
            text = content.decode("utf-8", errors="replace")
        else:
            raise HTTPException(400, "Only PDF, DOCX, and TXT files are supported.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Could not read {filename}: {exc}")
    if not text.strip():
        raise HTTPException(
            422,
            f"No readable text found in {filename}. If it is a scanned PDF, OCR is needed."
        )
    return text


def predict_category(text: str, top_k: int = 5) -> dict:
    if model is None:
        raise HTTPException(503, model_error or "Model is not available.")
    cleaned = clean_text(text)
    try:
        predicted = str(model.predict([cleaned])[0])
        result = {"predicted_category": predicted, "top_categories": []}
        if hasattr(model, "predict_proba"):
            probs = model.predict_proba([cleaned])[0]
            classes = model.named_steps["classifier"].classes_
            order = probs.argsort()[::-1][:top_k]
            result["top_categories"] = [
                {"category": str(classes[i]), "score": round(float(probs[i]), 4)}
                for i in order
            ]
        return result
    except Exception as exc:
        raise HTTPException(500, f"Prediction failed: {exc}")


def match_resume(resume_text: str, job_description: str) -> dict:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity

    resume_clean, job_clean = clean_text(resume_text), clean_text(job_description)
    try:
        vectorizer = TfidfVectorizer(
            stop_words="english", ngram_range=(1, 2), sublinear_tf=True
        )
        matrix = vectorizer.fit_transform([job_clean, resume_clean])
        similarity = float(cosine_similarity(matrix[0:1], matrix[1:2])[0, 0])
    except ValueError:
        similarity = 0.0

    job_skills, resume_skills = set(extract_skills(job_description)), set(extract_skills(resume_text))
    coverage = len(job_skills & resume_skills) / len(job_skills) if job_skills else 0.0
    score = (0.60 * similarity + 0.40 * coverage) * 100
    return {
        "match_score_heuristic": round(score, 2),
        "text_similarity_pct": round(similarity * 100, 2),
        "skill_coverage_pct": round(coverage * 100, 2),
        "matched_skills": sorted(job_skills & resume_skills),
        "missing_skills": sorted(job_skills - resume_skills),
        "resume_skills_detected": sorted(resume_skills),
        "job_skills_detected": sorted(job_skills),
        "note": "Heuristic score only; not a probability of job performance. Review resumes manually."
    }


@app.get("/")
def root():
    return {
        "message": "Resume Analyzer API is running",
        "model_loaded": model is not None,
        "model_status": "ready" if model is not None else model_error,
        "docs": "/docs"
    }


@app.get("/api/health")
def health():
    return {"status": "ok", "model_loaded": model is not None, "model_status": model_error}


@app.post("/api/predict")
def predict(request: TextPredictionRequest):
    return predict_category(request.text, request.top_k)


@app.post("/api/match")
def match(request: MatchRequest):
    return match_resume(request.resume_text, request.job_description)


@app.post("/api/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    job_description: str = Form(...),
):
    if not job_description.strip():
        raise HTTPException(400, "Job description is required.")
    if file.size and file.size > 10 * 1024 * 1024:
        raise HTTPException(413, "File exceeds 10 MB limit.")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "File exceeds 10 MB limit.")
    resume_text = extract_resume_text(file.filename or "resume.txt", content)
    category = None
    if model is not None:
        category = predict_category(resume_text, top_k=5)
    matching = match_resume(resume_text, job_description)
    return {
        "filename": file.filename,
        "characters_extracted": len(resume_text),
        "category_prediction": category,
        "matching": matching,
    }


@app.post("/api/rank")
async def rank_resumes(
    files: list[UploadFile] = File(...),
    job_description: str = Form(...),
):
    if not job_description.strip():
        raise HTTPException(400, "Job description is required.")
    if not files:
        raise HTTPException(400, "Upload at least one resume.")

    results = []
    for file in files:
        content = await file.read()
        if len(content) > 10 * 1024 * 1024:
            results.append({"filename": file.filename, "error": "File exceeds 10 MB limit."})
            continue
        try:
            resume_text = extract_resume_text(file.filename or "resume.txt", content)
            result = {
                "filename": file.filename,
                "characters_extracted": len(resume_text),
                **match_resume(resume_text, job_description),
            }
            if model is not None:
                result["category_prediction"] = predict_category(resume_text, top_k=3)
            results.append(result)
        except HTTPException as exc:
            results.append({"filename": file.filename, "error": exc.detail})

    results.sort(
        key=lambda row: row.get("match_score_heuristic", -1),
        reverse=True
    )
    return {
        "count": len(results),
        "results": results,
        "note": "Rankings are heuristic decision-support only; human review is required."
    }
